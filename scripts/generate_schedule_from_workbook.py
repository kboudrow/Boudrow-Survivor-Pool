from __future__ import annotations

import collections
import re
from datetime import date, datetime, time
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

WORKBOOK = Path(r"C:\Users\boudr\Downloads\2026_NFL_Schedule_Google_Sheets_AUDITED.xlsx")
OUTPUT = Path(__file__).resolve().parents[1] / "supabase" / "nfl-schedule-2026.sql"

MONTHS = {"Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12, "Jan": 1}
ET = ZoneInfo("America/New_York")
UTC = ZoneInfo("UTC")
WEEK_SUNDAYS = {
    1: "2026-09-13",
    2: "2026-09-20",
    3: "2026-09-27",
    4: "2026-10-04",
    5: "2026-10-11",
    6: "2026-10-18",
    7: "2026-10-25",
    8: "2026-11-01",
    9: "2026-11-08",
    10: "2026-11-15",
    11: "2026-11-22",
    12: "2026-11-29",
    13: "2026-12-06",
    14: "2026-12-13",
    15: "2026-12-20",
    16: "2026-12-27",
    17: "2027-01-03",
    18: "2027-01-10",
}


def sql(value: object) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def kickoff_utc(date_text: object, time_text: object, week: int) -> str:
    if not date_text or not time_text or "TBD" in str(date_text) or "TBD" in str(time_text):
        local_date = datetime.strptime(WEEK_SUNDAYS[week], "%Y-%m-%d").date()
        local_dt = datetime.combine(local_date, time(13, 0), ET)
        return local_dt.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")

    match = re.match(r"([A-Za-z]{3})\s+(\d+)", str(date_text))
    if not match:
        raise ValueError(f"Unsupported date value: {date_text!r}")

    month = MONTHS[match.group(1)]
    day = int(match.group(2))
    year = 2027 if month == 1 else 2026
    local_time = datetime.strptime(str(time_text).replace(" ET", "").strip(), "%I:%M %p").time()
    local_dt = datetime.combine(date(year, month, day), local_time, ET)
    return local_dt.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


def main() -> None:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True)
    sheet = workbook["2026 NFL Schedule"]
    rows = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        week = int(row[0])
        date_text = row[2]
        time_text = row[3]
        away = str(row[4]).strip()
        home = str(row[5]).strip()
        kickoff = kickoff_utc(date_text, time_text, week)
        event_id = f"audited-2026-w{week:02d}-{away}-{home}"
        rows.append((2026, week, kickoff, kickoff, home, away, "scheduled", event_id))

    week_counts = collections.Counter(row[1] for row in rows)
    team_counts: collections.Counter[str] = collections.Counter()
    week_teams: dict[int, list[str]] = {}
    for _, week, _, _, home, away, _, _ in rows:
        team_counts[home] += 1
        team_counts[away] += 1
        week_teams.setdefault(week, []).extend([home, away])

    double_booked = {
        week: [team for team, count in collections.Counter(teams).items() if count > 1]
        for week, teams in week_teams.items()
    }
    double_booked = {week: teams for week, teams in double_booked.items() if teams}

    if len(rows) != 272:
        raise SystemExit(f"Expected 272 games, found {len(rows)}")
    if min(team_counts.values()) != 17 or max(team_counts.values()) != 17:
        raise SystemExit(f"Expected every team to have 17 games, found {dict(team_counts)}")
    if double_booked:
        raise SystemExit(f"Teams double-booked in a week: {double_booked}")

    lines: list[str] = []
    lines.append("-- Generated from 2026_NFL_Schedule_Google_Sheets_AUDITED.xlsx.\n")
    lines.append("-- Source workbook audit: 272 games, all 32 teams have 17 games, no duplicate games, no same-week double-booked teams.\n")
    lines.append("-- Week 18 exact TV windows are TBD in the workbook; placeholders use Sunday 1 PM ET until official times are available.\n\n")
    lines.append("begin;\n\n")
    lines.append("delete from public.nfl_games\nwhere season = 2026;\n\n")
    lines.append("insert into public.season_weeks (season, week, week_sunday_date)\nvalues\n")
    lines.append(",\n".join(f"  (2026, {week}, '{sunday}')" for week, sunday in WEEK_SUNDAYS.items()))
    lines.append("\non conflict (season, week) do update\nset week_sunday_date = excluded.week_sunday_date;\n\n")
    lines.append(
        "insert into public.nfl_games (\n"
        "  season,\n"
        "  week,\n"
        "  game_time,\n"
        "  kickoff_at_utc,\n"
        "  home_team,\n"
        "  away_team,\n"
        "  status,\n"
        "  espn_event_id\n"
        ")\nvalues\n"
    )
    lines.append(
        ",\n".join(
            f"  ({season}, {week}, {sql(game_time)}, {sql(kickoff)}, {sql(home)}, {sql(away)}, {sql(status)}, {sql(event_id)})"
            for season, week, game_time, kickoff, home, away, status, event_id in rows
        )
    )
    lines.append(
        "\non conflict (espn_event_id) do update\n"
        "set\n"
        "  season = excluded.season,\n"
        "  week = excluded.week,\n"
        "  game_time = excluded.game_time,\n"
        "  kickoff_at_utc = excluded.kickoff_at_utc,\n"
        "  home_team = excluded.home_team,\n"
        "  away_team = excluded.away_team,\n"
        "  status = excluded.status;\n\n"
        "commit;\n"
    )

    OUTPUT.write_text("".join(lines), encoding="utf-8")
    print(f"wrote {len(rows)} games")
    print(dict(sorted(week_counts.items())))


if __name__ == "__main__":
    main()
